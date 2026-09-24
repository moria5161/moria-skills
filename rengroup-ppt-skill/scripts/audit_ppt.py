"""按任组演示文稿规范审计 PPTX。

内置模板默认执行严格样式检查；传入自定义模板时，固定样式差异默认降为警告，
避免把旧模板的数值规则误判成新版模板错误。使用 --strict-style 可强制严格模式。
"""

import argparse
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN


DEFAULT_TEMPLATE = Path(__file__).resolve().parents[1] / "assets" / "RenGroup-PPT-template.pptx"
BLACK = "000000"
WHITE = "FFFFFF"
EMU_PER_INCH = 914400
DATE_RE = re.compile(r"(?:20\d{2}[./年-]\d{1,2}|日期|研究日期|汇报时间)")


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument(
        "--template",
        type=Path,
        default=DEFAULT_TEMPLATE,
        help="用于读取标题色和结论框颜色的模板。",
    )
    parser.add_argument(
        "--strict-style",
        action="store_true",
        help="即使使用自定义模板，也按内置任组固定样式把差异视为错误。",
    )
    return parser.parse_args()


def configure_console():
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure:
            reconfigure(encoding="utf-8", errors="replace")


def template_colors(template_path):
    with ZipFile(template_path) as archive:
        master = archive.read("ppt/slideMasters/slideMaster1.xml").decode("utf-8")
        title_match = re.search(
            r'<p:titleStyle>.*?<a:solidFill>\s*<a:srgbClr val="([0-9A-Fa-f]{6})"',
            master,
            re.DOTALL,
        )
        title_color = title_match.group(1).upper() if title_match else "800000"

        conclusion_color = None
        for name in archive.namelist():
            if not re.fullmatch(r"ppt/slides/slide\d+\.xml", name):
                continue
            slide_xml = archive.read(name).decode("utf-8")
            frame_match = re.search(
                r'<a:prstGeom prst="roundRect".*?<a:ln[^>]*>.*?'
                r'<a:solidFill>\s*<a:srgbClr val="([0-9A-Fa-f]{6})"',
                slide_xml,
                re.DOTALL,
            )
            if frame_match:
                conclusion_color = frame_match.group(1).upper()
                break
    return title_color, conclusion_color or "C00000"


def explicit_rgb(color_format):
    try:
        value = color_format.rgb
        return str(value).upper() if value is not None else None
    except (AttributeError, TypeError):
        return None


def explicit_fill_rgb(fill_format):
    try:
        return explicit_rgb(fill_format.fore_color)
    except (AttributeError, TypeError):
        return None


def is_title(shape):
    if not getattr(shape, "is_placeholder", False):
        return False
    try:
        return shape.placeholder_format.type in (
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE,
        )
    except ValueError:
        return False


def iter_runs(shape):
    if not getattr(shape, "has_text_frame", False):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                yield paragraph, run


def contains_cjk(text):
    return any(
        "\u3400" <= char <= "\u4dbf"
        or "\u4e00" <= char <= "\u9fff"
        or "\uf900" <= char <= "\ufaff"
        for char in text
    )


def within(inner, outer):
    return (
        inner.left >= outer.left
        and inner.top >= outer.top
        and inner.left + inner.width <= outer.left + outer.width
        and inner.top + inner.height <= outer.top + outer.height
    )


def style_issue(message, errors, warnings, strict_style):
    (errors if strict_style else warnings).append(message)


def audit(path, title_color, conclusion_color, strict_style):
    prs = Presentation(path)
    errors, warnings = [], []
    sw, sh = prs.slide_width, prs.slide_height

    for slide_no, slide in enumerate(prs.slides, 1):
        titles = [shape for shape in slide.shapes if is_title(shape)]
        if slide_no > 1:
            if not titles:
                style_issue(
                    f"第 {slide_no} 页：未找到标题占位符。",
                    errors,
                    warnings,
                    strict_style,
                )
            for title in titles:
                center_offset = abs((title.left + title.width / 2) - sw / 2) / EMU_PER_INCH
                if center_offset > 0.25 or title.top > 1.15 * EMU_PER_INCH:
                    style_issue(
                        f"第 {slide_no} 页：标题未位于内置模板的顶部居中区域。",
                        errors,
                        warnings,
                        strict_style,
                    )
                for paragraph in title.text_frame.paragraphs:
                    if paragraph.text.strip() and paragraph.alignment != PP_ALIGN.CENTER:
                        warnings.append(f"第 {slide_no} 页：标题未显式设置为居中对齐。")
                    for run in paragraph.runs:
                        if not run.text.strip():
                            continue
                        color = explicit_rgb(run.font.color)
                        if color is None:
                            warnings.append(
                                f"第 {slide_no} 页：标题颜色来自继承，请确认模板最终解析为 #{title_color}。"
                            )
                        elif color != title_color:
                            style_issue(
                                f"第 {slide_no} 页：标题颜色为 #{color}，模板参考为 #{title_color}。",
                                errors,
                                warnings,
                                strict_style,
                            )

        red_frames = []
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                line_color = explicit_rgb(shape.line.color)
                if line_color == conclusion_color:
                    red_frames.append(shape)
                    low_enough = shape.top >= sh * 0.62 or (
                        shape.left >= sw * 0.48 and shape.top >= sh * 0.45
                    )
                    if not low_enough:
                        style_issue(
                            f"第 {slide_no} 页：结论框不在内置模板常用的底部或右下区域。",
                            errors,
                            warnings,
                            strict_style,
                        )
                    if shape.fill.type is not None:
                        fill_color = explicit_fill_rgb(shape.fill)
                        if fill_color not in (None, WHITE):
                            style_issue(
                                f"第 {slide_no} 页：结论框存在非白色填充。",
                                errors,
                                warnings,
                                strict_style,
                            )

        conclusion_text = [
            shape
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and "结论" in shape.text
        ]
        if conclusion_text and not red_frames:
            style_issue(
                f"第 {slide_no} 页：检测到结论文字，但没有模板色 #{conclusion_color} 的结论框。",
                errors,
                warnings,
                strict_style,
            )
        for text_shape in conclusion_text:
            if not any(within(text_shape, frame) for frame in red_frames):
                warnings.append(f"第 {slide_no} 页：结论文字未完全位于结论框内。")
            for _, run in iter_runs(text_shape):
                color = explicit_rgb(run.font.color)
                if color not in (None, BLACK):
                    style_issue(
                        f"第 {slide_no} 页：结论文字不是黑色。",
                        errors,
                        warnings,
                        strict_style,
                    )
                if run.font.size and not 20 <= run.font.size.pt <= 24:
                    style_issue(
                        f"第 {slide_no} 页：结论文字为 {run.font.size.pt:g} pt，内置规范参考 20–24 pt。",
                        errors,
                        warnings,
                        strict_style,
                    )

        for shape in slide.shapes:
            text = getattr(shape, "text", "").strip()
            if slide_no > 1 and text and DATE_RE.search(text):
                if shape.left < sw * 0.45 and shape.top > sh * 0.72:
                    errors.append(f"第 {slide_no} 页：内容页左下角不允许出现日期或时间标记。")

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    if shape.image.ext.lower() not in ("svg", "emf", "wmf"):
                        px_w, px_h = shape.image.size
                        ppi = min(
                            px_w / (shape.width / EMU_PER_INCH),
                            px_h / (shape.height / EMU_PER_INCH),
                        )
                        if ppi < 150:
                            warnings.append(
                                f'第 {slide_no} 页：图片“{shape.name}”在当前显示尺寸下仅有 {ppi:.0f} ppi。'
                            )
                except (AttributeError, ZeroDivisionError):
                    pass

            if getattr(shape, "has_table", False) and slide_no > 1:
                for row in shape.table.rows:
                    for cell in row.cells:
                        fill_color = (
                            explicit_fill_rgb(cell.fill)
                            if cell.fill.type is not None
                            else None
                        )
                        if fill_color not in (None, WHITE):
                            style_issue(
                                f"第 {slide_no} 页：表格单元格存在彩色填充。",
                                errors,
                                warnings,
                                strict_style,
                            )
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if (
                                    run.text.strip()
                                    and run.font.size
                                    and not 14 <= run.font.size.pt <= 16
                                ):
                                    style_issue(
                                        f"第 {slide_no} 页：表格文字为 {run.font.size.pt:g} pt，内置规范参考 14–16 pt。",
                                        errors,
                                        warnings,
                                        strict_style,
                                    )

            if is_title(shape) or slide_no == 1:
                continue

            for _, run in iter_runs(shape):
                size = run.font.size.pt if run.font.size else None
                if size is not None and size < 12:
                    style_issue(
                        f'第 {slide_no} 页：文字“{run.text[:24]}”小于 12 pt。',
                        errors,
                        warnings,
                        strict_style,
                    )
                elif size is not None and 16 < size < 20:
                    warnings.append(
                        f'第 {slide_no} 页：文字“{run.text[:24]}”为 {size:g} pt，请确认其角色与可读性。'
                    )
                elif size is not None and size > 24:
                    warnings.append(
                        f'第 {slide_no} 页：非标题文字“{run.text[:24]}”超过 24 pt。'
                    )

                has_cn = contains_cjk(run.text)
                has_en = bool(re.search(r"[A-Za-z]", run.text))
                font = run.font.name
                if font is None:
                    warnings.append(
                        f'第 {slide_no} 页：文字“{run.text[:24]}”的字体来自继承，请在最终渲染中确认。'
                    )
                else:
                    if has_cn and font not in ("黑体", "SimHei"):
                        style_issue(
                            f'第 {slide_no} 页：中文文字使用“{font}”，内置组内规范为黑体。',
                            errors,
                            warnings,
                            strict_style,
                        )
                    if has_en and font != "Arial":
                        style_issue(
                            f'第 {slide_no} 页：英文文字使用“{font}”，内置组内规范为 Arial。',
                            errors,
                            warnings,
                            strict_style,
                        )

            if getattr(shape, "has_chart", False):
                chart = shape.chart
                for axis_name in ("category_axis", "value_axis"):
                    try:
                        axis = getattr(chart, axis_name)
                        size = axis.tick_labels.font.size
                        if size and size.pt < 12:
                            style_issue(
                                f"第 {slide_no} 页：{axis_name} 刻度标签小于 12 pt。",
                                errors,
                                warnings,
                                strict_style,
                            )
                    except (AttributeError, ValueError):
                        pass
                try:
                    size = chart.legend.font.size
                    if size and not 14 <= size.pt <= 16:
                        style_issue(
                            f"第 {slide_no} 页：图例为 {size.pt:g} pt，内置规范参考 14–16 pt。",
                            errors,
                            warnings,
                            strict_style,
                        )
                except (AttributeError, ValueError):
                    pass

    return errors, warnings


def main():
    configure_console()
    args = parse_args()
    title_color, conclusion_color = template_colors(args.template)

    try:
        using_builtin = args.template.resolve() == DEFAULT_TEMPLATE.resolve()
    except OSError:
        using_builtin = False
    strict_style = using_builtin or args.strict_style

    mode = "严格组内样式" if strict_style else "自定义模板兼容"
    print(f"审计模式：{mode}")
    print(f"模板颜色：标题 #{title_color}，结论框 #{conclusion_color}")

    errors, warnings = audit(
        args.pptx,
        title_color,
        conclusion_color,
        strict_style,
    )
    for message in warnings:
        print("警告：", message)
    for message in errors:
        print("错误：", message)

    print(f"审计完成：{len(errors)} 个错误，{len(warnings)} 个警告。")
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()
